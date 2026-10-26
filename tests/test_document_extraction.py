"""Black-box contract tests for HTML-to-Markdown conversion, Office/TCVN3 recovery, and PDF/OCR extraction."""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from .conftest import make_native_pdf, make_scanned_pdf


class TestHtmlToMarkdownConversion:
    def test_scoped_html_with_tables_and_relative_links_renders_deterministic_markdown(
        self,
    ) -> None:
        """HTML within `content_xpath` converts headings, relative links, and tables into Markdown."""
        from wlh.extraction.markdown import html_to_markdown

        html = """
        <html>
          <body>
            <main id="legal-body">
              <h1>Act No. 42/2026</h1>
              <p>See <a href="/docs/annex1.pdf">Annex 1</a> for details.</p>
              <table>
                <thead>
                  <tr><th>Article</th><th>Subject</th></tr>
                </thead>
                <tbody>
                  <tr><td>Article 1</td><td>Scope of Application</td></tr>
                  <tr><td>Article 2</td><td>Definitions</td></tr>
                </tbody>
              </table>
            </main>
          </body>
        </html>
        """
        md = html_to_markdown(
            html,
            base_url="https://law.example.org/acts/42",
            content_xpath="//main[@id='legal-body']",
        )
        assert "# Act No. 42/2026" in md
        assert "[Annex 1](https://law.example.org/docs/annex1.pdf)" in md
        assert "| Article | Subject |" in md
        assert "| Article 1 | Scope of Application |" in md

    def test_html_with_ui_chrome_font_controls_print_back_buttons_and_share_bars_renders_clean_markdown(
        self,
    ) -> None:
        """HTML containing UI chrome (font widgets, print/back buttons, share bars, anchors without href, adjacent bolds) renders clean Markdown."""
        from wlh.extraction.markdown import html_to_markdown

        html = """
        <html>
          <body>
            <div class="extend_font fl">【字体：<a href="javascript:fontZoom(26)">大</a> <a href="javascript:fontZoom(20)">中</a> <a href="javascript:fontZoom(14)">小</a>】</div>
            <div>2026-06-30 字体 [大 中 小]</div>
            <div class="eyd-page-header__share">
              <span>Share:</span>
              <ul>
                <li><a href="#" aria-label="Facebook"><span class="eyd-icon-facebook"></span></a></li>
                <li><a href="#" aria-label="X"><span class="eyd-icon-twitter"></span></a></li>
              </ul>
            </div>
            <div class="eyd-page-anchors">
              <span>On this page: </span>
              <nav><ul><li><a href="#sec1">Section 1</a></li></ul></nav>
            </div>
            <button type="button">Print this page</button>
            <button type="button">Show contents</button>
            <a id="main-content"></a>
            <a href="#NOT-2026-46" class="xref"></a>
            <p><b><span>第</span></b><b><span lang="EN-US">1</span></b><b><span>号修改单的通知</span></b></p>
            <p><a name="p1" id="p1">1.</a> For patients, the dose is <a name="_Hlk234317469">100 mg</a> daily.</p>
            <p>See <i>Lui Sai-yu </i><b>at </b><b>paragraphs </b>58.</p>
            <div class="document-tip target-hidden"><span>Tải biểu mẫu</span></div>
            <span class="tooltip-button target-hidden"><span>Đang theo dõi</span></span>
            <div class="zrbj">[责任编辑： 贾欣然]</div>
            <div id="xgxw" style="display: none;"><h3>相关新闻</h3></div>
            <a href="#" objid="135" onclick="javascript:window.print()">【打印】</a>
            <a href="#" objid="135" onclick="goSpecialURL()">【返回】</a>
            <section class="eyd-related-pages">
              <h2>Related Content</h2>
              <article><a href="/other">Related Guide</a></article>
            </section>
          </body>
        </html>
        """
        md = html_to_markdown(
            html,
            base_url="https://www.safe.gov.cn/www/illegal/listQuery?saffno=22d58c4edfbf401595848a52a3f91db6",
        )
        for forbidden in (
            "【字体：",
            "字体 [大 中 小]",
            "相关新闻",
            "责任编辑",
            "【打印】",
            "【返回】",
            "Print this page",
            "Show contents",
            "Share:",
            "On this page:",
            "Related Content",
            "Đang theo dõi",
            "Tải biểu mẫu",
            "[](https://",
            "#)",
            "[1.](https://",
            "[100 mg](https://",
            "****",
        ):
            assert forbidden not in md, f"Unexpected noise {forbidden!r} found in Markdown:\n{md}"

        assert "1. For patients, the dose is 100 mg daily." in md
        assert "**第1号修改单的通知**" in md
        assert "*Lui Sai-yu* **at** **paragraphs** 58." in md

    def test_irs_bulletin_html_with_left_toc_column_and_named_anchors_excludes_front_toc_noise(
        self,
    ) -> None:
        """IRS IRB two-column layout strips the `col-md-4` table-of-contents sidebar and empty anchor links while keeping `col-md-8` body text."""
        from wlh.extraction.markdown import html_to_markdown

        html = (
            '<div id="block-pup-irs-barrio-content">'
            '<div class="field--name-body">'
            '<div class="row">'
            '<div class="col-md-4">'
            "<ul>"
            '<li><a href="#id1">HIGHLIGHTS OF THIS ISSUE</a>'
            '<ul><li><a href="#HI1">INCOME TAX</a></li></ul>'
            "</li>"
            '<li><a href="#id3"></a>'
            '<ul><li><a href="#id5">The IRS Mission</a></li></ul>'
            "</li>"
            "</ul>"
            "</div>"
            '<div class="col-md-8">'
            '<div class="book">'
            '<div class="titlepage">'
            "<h1>Internal Revenue Bulletin: 2026-35</h1>"
            '<p class="pubdate">August 24, 2026</p>'
            "</div>"
            '<div class="part">'
            '<h1><a name="id1"></a>HIGHLIGHTS OF THIS ISSUE</h1>'
            "<p>Substantive bulletin text on Section 43 inflation adjustment.</p>"
            "</div>"
            "</div>"
            "</div>"
            "</div>"
            "</div>"
            "</div>"
        )
        md = html_to_markdown(
            html,
            base_url="https://www.irs.gov/irb/2026-35_IRB",
            content_xpath='//div[@id="block-pup-irs-barrio-content"]',
        )
        for forbidden in (
            "[HIGHLIGHTS OF THIS ISSUE](",
            "[INCOME TAX](",
            "[The IRS Mission](",
            "[](https://www.irs.gov/irb/2026-35_IRB",
        ):
            assert forbidden not in md, f"Unexpected TOC/anchor noise {forbidden!r} in Markdown:\n{md}"

        assert "# Internal Revenue Bulletin: 2026-35" in md
        assert "Substantive bulletin text on Section 43 inflation adjustment." in md


class TestLegacyVietnameseAndDisguisedOfficeExtraction:
    def test_mixed_unicode_and_tcvn3_lines_recover_vietnamese_diacritics(self) -> None:
        """Line-level TCVN3 recovery preserves existing Unicode headers while converting TCVN3 lines."""
        from wlh.extraction.vi_tcvn3 import normalize_detected_tcvn3

        unicode_header = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
        tcvn3_line = (
            "Lu\xcbt Doanh nghi\xd6p s\xe8 59/2020/QH14 \xa7i\xd2u 1 Ph\xb5m vi \xaei\xd2u ch\xd8nh"
        )
        mixed_text = f"{unicode_header}\n{tcvn3_line}\n"

        recovered, changed = normalize_detected_tcvn3(mixed_text)
        assert changed is True
        assert unicode_header in recovered
        assert "Luật Doanh nghiệp" in recovered

    def test_disguised_html_and_rtf_doc_files_detect_actual_format_and_extract_text(
        self, tmp_path: Path
    ) -> None:
        """Files named `.doc` that actually contain HTML or TCVN3 RTF are sniffed and extracted accurately."""
        from wlh.extraction.office import extract_office

        disguised_html_doc = tmp_path / "disguised_html.doc"
        disguised_html_doc.write_bytes(
            b"<html><body><h1>Circular 01/2026/TT-BTC</h1><p>Tax regulation body text.</p></body></html>"
        )
        md_html, info_html = extract_office(disguised_html_doc, kind="doc")
        assert "Circular 01/2026/TT-BTC" in md_html
        assert info_html["detected_format"] == "html"

        disguised_rtf_doc = tmp_path / "disguised_rtf.doc"
        rtf_bytes = (
            b"{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 .VnTime;}}\n"
            b"\\f0\\fs24 Lu\\'cbt Doanh nghi\\'d6p s\\'e8 59/2020/QH14 \\'a7i\\'d2u 1 Ph\\'b5m vi \\'aei\\'d2u ch\\'d8nh\\par}"
        )
        disguised_rtf_doc.write_bytes(rtf_bytes)
        md_rtf, info_rtf = extract_office(disguised_rtf_doc, kind="doc")
        assert "Luật Doanh nghiệp" in md_rtf
        assert info_rtf["detected_format"] == "rtf"

    def test_scanned_image_docx_with_only_signature_text_falls_back_to_ocr_or_rejects_stub(
        self, tmp_path: Path
    ) -> None:
        """A `.docx` containing only a 56-char signature text block and an embedded scan image in `word/media/image1.png` triggers OCR fallback or rejects the stub."""
        import inspect
        import zipfile
        from typing import Any

        import fitz  # type: ignore[import-untyped]
        from docx import Document  # type: ignore[import-untyped]
        from docx.shared import Inches  # type: ignore[import-untyped]

        from wlh.core import WlhError
        from wlh.extraction.office import extract_office

        fitz_mod: Any = fitz
        img_doc = fitz_mod.open()
        img_page = img_doc.new_page(width=595, height=420)
        img_page.insert_textbox(
            fitz_mod.Rect(30, 40, 565, 380),
            "DECREE NO 325 FULL STATUTORY TEXT ON PUBLIC ADMINISTRATION\n"
            "Article 1. All public agencies shall digitize administrative records.",
            fontsize=18,
            fontname="helv",
        )
        img_path = tmp_path / "scan_body.png"
        img_page.get_pixmap(dpi=150).save(str(img_path))
        img_doc.close()

        sig_text = "TM. CHINH PHU KT. THU TUONG PHO THU TUONG Phan Van Giang"
        assert len(sig_text) == 56

        docx_path = tmp_path / "scanned_sig.docx"
        word_doc: Any = Document()
        word_doc.add_paragraph(sig_text)
        word_doc.add_picture(str(img_path), width=Inches(6.0))
        word_doc.save(str(docx_path))

        with zipfile.ZipFile(docx_path) as zf:
            names = set(zf.namelist())
            assert "word/document.xml" in names
            assert "word/media/image1.png" in names

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 60,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        sig = inspect.signature(extract_office)
        extract_fn: Any = extract_office
        try:
            if "ocr" in sig.parameters:
                md_docx, info_docx = extract_fn(
                    docx_path, kind="docx", ocr=ocr_cfg, script_hint="latin"
                )
            else:
                md_docx, info_docx = extract_office(docx_path, kind="docx")
        except WlhError as exc:
            assert exc.code in {
                "quality_too_short",
                "quality_check_failed",
                "docx_text_unusable",
            }
        else:
            assert info_docx.get("extractor") not in {"python_docx", "docx_xml"}
            assert "DECREE NO 325" in md_docx.upper()


class TestPdfExtractionAndOcrFallback:
    def test_digital_born_pdf_extracts_text_directly_without_invoking_ocr(
        self, tmp_path: Path
    ) -> None:
        """A digital-born PDF with selectable text is extracted directly without triggering OCR."""
        from wlh.extraction.pdf_scan import analyze, extract_pdf

        native_pdf = make_native_pdf(
            tmp_path / "native.pdf",
            [
                "FEDERAL REGISTER VOL. 91 NO. 10\nPresidential Document 2026-00101 on Open Legal Data.",
                "Section 2. All agencies shall publish machine-readable statutory instruments.",
            ],
        )
        analysis, _ = analyze(native_pdf)
        assert analysis["page_count"] == 2
        assert analysis["suspected_scan_pages"] == []

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 60,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        md_native, info_native, _, parsers_native = extract_pdf(native_pdf, ocr_cfg, "latin")
        assert "Presidential Document 2026-00101" in md_native
        assert info_native["method"] == "native"
        assert info_native["ocr"] is None
        assert "ocr" not in parsers_native

    def test_scanned_pdf_with_unreachable_primary_ocr_falls_back_to_local_ocr(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A scanned PDF automatically falls back to local OCR when the remote OCR endpoint is unreachable."""
        from wlh.extraction.pdf_scan import analyze, extract_pdf

        monkeypatch.setenv("WLH_BR_OCR_URL", "http://127.0.0.1:59999/api/v1")
        scanned_pdf = make_scanned_pdf(
            tmp_path / "scanned.pdf",
            [
                "SUPREME COURT OF THE UNITED STATES\nOPINION NO. 25-999\nJudgment Affirmed on October 20, 2026."
            ],
        )
        scan_analysis, _ = analyze(scanned_pdf)
        assert scan_analysis["suspected_scan_pages"] == [1]

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 60,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        t0 = time.monotonic()
        md_scan, info_scan, _, parsers_scan = extract_pdf(scanned_pdf, ocr_cfg, "latin")
        elapsed = time.monotonic() - t0

        assert elapsed < 15.0, f"Fallback took too long: {elapsed:.2f}s"
        assert "SUPREME COURT" in md_scan.upper()
        assert info_scan["method"] == "ocr"
        assert info_scan["ocr"] is not None
        assert info_scan["ocr"]["engine"] == "tesseract"
        assert info_scan["ocr"].get("fallback_from") == "br-ocr-v1"
        assert "ocr" in parsers_scan

    def test_scanned_pdf_with_sparse_signature_text_layer_triggers_ocr_pipeline(
        self, tmp_path: Path
    ) -> None:
        """A scanned PDF with a small digital signature block text layer triggers the OCR pipeline rather than being accepted as pure text."""
        from typing import Any

        import fitz  # type: ignore[import-untyped]

        from wlh.extraction.pdf_scan import analyze, extract_pdf

        fitz_mod: Any = fitz
        temp_doc = fitz_mod.open()
        p = temp_doc.new_page(width=595, height=842)
        p.insert_textbox(
            fitz_mod.Rect(50, 80, 545, 760),
            "OFFICIAL GAZETTE DECREE NO 2026-88\nFull statutory text of the regulation governing maritime security.",
            fontsize=18,
            fontname="helv",
        )
        pix = p.get_pixmap(dpi=150)
        img_bytes = pix.tobytes("png")
        temp_doc.close()

        sig_doc = fitz_mod.open()
        p2 = sig_doc.new_page(width=595, height=842)
        p2.insert_image(fitz_mod.Rect(0, 0, 595, 842), stream=img_bytes)
        p2.insert_textbox(
            fitz_mod.Rect(50, 750, 500, 800),
            "Signer: Ministry Office\nSigned at: 2026-10-23 10:00:00\nCertificate ID: 8849-1102",
            fontsize=10,
            fontname="helv",
        )
        pdf_path = tmp_path / "sig_scan.pdf"
        sig_doc.save(str(pdf_path))
        sig_doc.close()

        analysis, _ = analyze(pdf_path)
        assert analysis["suspected_scan_pages"] == [1]

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 30,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        md, info, _, parsers = extract_pdf(pdf_path, ocr_cfg, "latin")
        assert "OFFICIAL GAZETTE" in md.upper()
        assert info["method"] == "ocr"
        assert "ocr" in parsers

    def test_digital_born_pdf_with_cover_image_does_not_trigger_scan_classification(
        self, tmp_path: Path
    ) -> None:
        """A digital-born PDF whose cover page has a large image but has normal text layers is not misclassified as a scan."""
        from typing import Any

        import fitz  # type: ignore[import-untyped]

        from wlh.extraction.pdf_scan import analyze, extract_pdf

        fitz_mod: Any = fitz
        doc = fitz_mod.open()

        cover_img = fitz_mod.open()
        cp = cover_img.new_page(width=595, height=842)
        cp.insert_textbox(
            fitz_mod.Rect(100, 200, 500, 400),
            "MINISTERIAL REPORT 2026\nANNUAL LEGAL GAZETTE",
            fontsize=24,
            fontname="helv",
        )
        pix = cp.get_pixmap(dpi=150)
        img_bytes = pix.tobytes("png")
        cover_img.close()

        p1 = doc.new_page(width=595, height=842)
        p1.insert_image(fitz_mod.Rect(0, 0, 595, 842), stream=img_bytes)
        p1.insert_textbox(
            fitz_mod.Rect(50, 750, 500, 800),
            "Official Publication of the Department of Legal Affairs",
            fontsize=12,
            fontname="helv",
        )

        p2 = doc.new_page(width=595, height=842)
        p2.insert_textbox(
            fitz_mod.Rect(50, 80, 545, 760),
            "Section 1. Purpose and Scope of this statutory regulation.\n"
            "This statutory instrument sets forth the standards for open government data compliance.\n"
            "All public departments shall prepare annual audit reviews and maintain machine-readable archives.\n"
            "The Director shall inspect all registries and certify the accuracy of public gazette records.\n"
            "Any inconsistencies between electronic representations and certified source documents must be recorded.",
            fontsize=14,
            fontname="helv",
        )
        pdf_path = tmp_path / "cover_image_native.pdf"
        doc.save(str(pdf_path))
        doc.close()

        analysis, _ = analyze(pdf_path)
        assert analysis["page_count"] == 2
        assert analysis["suspected_scan_pages"] == [], "Cover page was wrongly suspected as scan!"

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 30,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        md, info, _, parsers = extract_pdf(pdf_path, ocr_cfg, "latin")
        assert "ocr" not in parsers
        assert info["method"] == "native"
        assert "statutory regulation" in md

    def test_scanned_pdf_with_running_header_and_cropped_scan_image_triggers_ocr(
        self, tmp_path: Path
    ) -> None:
        """A 2-page gazette PDF with a ~50-60 char running header and a ~60% coverage scanned body image on each page triggers OCR."""
        from typing import Any

        import fitz  # type: ignore[import-untyped]

        from wlh.extraction.pdf_scan import analyze, extract_pdf

        fitz_mod: Any = fitz
        temp_doc = fitz_mod.open()
        tp = temp_doc.new_page(width=515, height=520)
        tp.insert_textbox(
            fitz_mod.Rect(20, 40, 495, 480),
            "GOVERNMENT RESOLUTION NO 325 ON DIGITAL INFRASTRUCTURE\n"
            "Article 1. National digital infrastructure standards and public data governance.",
            fontsize=18,
            fontname="helv",
        )
        img_bytes = tp.get_pixmap(dpi=150).tobytes("png")
        temp_doc.close()

        pdf_doc = fitz_mod.open()
        for page_no in (121, 122):
            p = pdf_doc.new_page(width=595, height=842)
            p.insert_textbox(
                fitz_mod.Rect(40, 25, 555, 80),
                f"CONG BAO / So 560 / Ngay 23-10-2026   Trang {page_no}",
                fontsize=11,
                fontname="helv",
            )
            p.insert_image(fitz_mod.Rect(40, 100, 555, 620), stream=img_bytes)
        pdf_path = tmp_path / "gazette_running_header_scan.pdf"
        pdf_doc.save(str(pdf_path))
        pdf_doc.close()

        analysis, _ = analyze(pdf_path)
        assert analysis["suspected_scan_pages"] == [1, 2]

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 60,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        md, info, _, parsers = extract_pdf(pdf_path, ocr_cfg, "latin")
        assert info["method"] == "ocr"
        assert "ocr" in parsers
        assert "RESOLUTION NO 325" in md.upper() or "DIGITAL INFRASTRUCTURE" in md.upper()
