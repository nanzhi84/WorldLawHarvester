"""Legacy Vietnamese text and disguised Office document extraction contracts."""

from __future__ import annotations

from pathlib import Path


class TestLegacyVietnameseAndDisguisedOfficeExtraction:
    def test_mixed_unicode_and_tcvn3_lines_recover_vietnamese_diacritics(self) -> None:
        """Line-level TCVN3 recovery preserves existing Unicode headers while converting TCVN3 lines."""
        from wlh.extraction.vi_tcvn3 import normalize_detected_tcvn3

        unicode_header = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
        tcvn3_line = (
            "Lu\xcbt Doanh nghi\xd6p s\xe8 59/2020/QH14 \xa7i\xd2u 1 Ph\xb5m vi \xaei\xd2u ch\xd8nh"
        )
        mixed_text = f"{unicode_header}\n{tcvn3_line}\n"

        recovered, changed = normalize_detected_tcvn3(mixed_text)
        assert changed is True
        assert unicode_header in recovered
        assert "Luật Doanh nghiệp" in recovered

    def test_disguised_html_and_rtf_doc_files_detect_actual_format_and_extract_text(
        self, tmp_path: Path
    ) -> None:
        """Files named `.doc` that actually contain HTML or TCVN3 RTF are sniffed and extracted accurately."""
        from wlh.extraction.office import extract_office

        disguised_html_doc = tmp_path / "disguised_html.doc"
        disguised_html_doc.write_bytes(
            b"<html><body><h1>Circular 01/2026/TT-BTC</h1><p>Tax regulation body text.</p></body></html>"
        )
        md_html, info_html = extract_office(disguised_html_doc, kind="doc")
        assert "Circular 01/2026/TT-BTC" in md_html
        assert info_html["detected_format"] == "html"

        disguised_rtf_doc = tmp_path / "disguised_rtf.doc"
        rtf_bytes = (
            b"{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 .VnTime;}}\n"
            b"\\f0\\fs24 Lu\\'cbt Doanh nghi\\'d6p s\\'e8 59/2020/QH14 \\'a7i\\'d2u 1 Ph\\'b5m vi \\'aei\\'d2u ch\\'d8nh\\par}"
        )
        disguised_rtf_doc.write_bytes(rtf_bytes)
        md_rtf, info_rtf = extract_office(disguised_rtf_doc, kind="doc")
        assert "Luật Doanh nghiệp" in md_rtf
        assert info_rtf["detected_format"] == "rtf"

    def test_scanned_image_docx_with_only_signature_text_falls_back_to_ocr_or_rejects_stub(
        self, tmp_path: Path
    ) -> None:
        """A `.docx` containing only a 56-char signature text block and an embedded scan image in `word/media/image1.png` triggers OCR fallback or rejects the stub."""
        import inspect
        import zipfile
        from typing import Any

        import fitz  # type: ignore[import-untyped]
        from docx import Document  # type: ignore[import-untyped]
        from docx.shared import Inches  # type: ignore[import-untyped]

        from wlh.core import WlhError
        from wlh.extraction.office import extract_office

        fitz_mod: Any = fitz
        img_doc = fitz_mod.open()
        img_page = img_doc.new_page(width=595, height=420)
        img_page.insert_textbox(
            fitz_mod.Rect(30, 40, 565, 380),
            "DECREE NO 325 FULL STATUTORY TEXT ON PUBLIC ADMINISTRATION\n"
            "Article 1. All public agencies shall digitize administrative records.",
            fontsize=18,
            fontname="helv",
        )
        img_path = tmp_path / "scan_body.png"
        img_page.get_pixmap(dpi=150).save(str(img_path))
        img_doc.close()

        sig_text = "TM. CHINH PHU KT. THU TUONG PHO THU TUONG Phan Van Giang"
        assert len(sig_text) == 56

        docx_path = tmp_path / "scanned_sig.docx"
        word_doc: Any = Document()
        word_doc.add_paragraph(sig_text)
        word_doc.add_picture(str(img_path), width=Inches(6.0))
        word_doc.save(str(docx_path))

        with zipfile.ZipFile(docx_path) as zf:
            names = set(zf.namelist())
            assert "word/document.xml" in names
            assert "word/media/image1.png" in names

        ocr_cfg = {
            "engine": "br-ocr-v1",
            "languages": "eng",
            "model": "br-ocr-v1",
            "timeout_seconds": 60,
            "poll_seconds": 0.2,
            "max_pages": 50,
            "dpi": 150,
            "max_page_pixels": 20000000,
            "max_result_bytes": 16777216,
        }
        sig = inspect.signature(extract_office)
        extract_fn: Any = extract_office
        try:
            if "ocr" in sig.parameters:
                md_docx, info_docx = extract_fn(
                    docx_path, kind="docx", ocr=ocr_cfg, script_hint="latin"
                )
            else:
                md_docx, info_docx = extract_office(docx_path, kind="docx")
        except WlhError as exc:
            assert exc.code in {
                "quality_too_short",
                "quality_check_failed",
                "docx_text_unusable",
            }
        else:
            assert info_docx.get("extractor") not in {"python_docx", "docx_xml"}
            assert "DECREE NO 325" in md_docx.upper()
